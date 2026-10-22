"""Pool of Apify API keys: one key per calendar day, automatic failover, and a usage log.

Rotation: each day the active key that was used least recently is picked and kept for the whole day
(so ~30 keys cover a month). If it turns out invalid or out of credit, it is marked exhausted for the
day and the next key takes over. With an empty pool, APIFY_TOKEN from .env is used as before.
"""

import csv
import io
import logging
import re
from dataclasses import dataclass
from datetime import UTC, date, datetime
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.apify.client import ApifyClient, ApifyError, summarize_limits
from app.config import get_settings
from app.logging_setup import redact
from app.models import ApifyKey, ApifyKeyDay, ApifyKeyUsage
from app.security import SecretError, decrypt, encrypt, fingerprint, mask

log = logging.getLogger("apify_keys")

TOKEN_RE = re.compile(r"apify_api_[A-Za-z0-9]{20,}")
ENV_LABEL = "APIFY_TOKEN (.env)"


class KeyPoolError(Exception):
    pass


@dataclass
class ResolvedKey:
    key_id: int | None  # None = the .env APIFY_TOKEN
    label: str
    token: str

    @property
    def masked(self) -> str:
        return mask(self.token[-4:])


def today() -> date:
    return datetime.now(ZoneInfo(get_settings().key_rotation_tz)).date()


def _env_token() -> str | None:
    t = get_settings().apify_token
    return t.get_secret_value() if t and t.get_secret_value() else None


def key_out(k: ApifyKey, uses: int = 0, today_id: int | None = None) -> dict:
    d = today()
    return {
        "id": k.id, "label": k.label, "masked": mask(k.last4), "active": k.active,
        "exhausted_today": k.exhausted_on == d, "last_error": k.last_error,
        "last_used_on": k.last_used_on, "created_at": k.created_at, "uses": uses,
        "is_today": k.id == today_id,
    }


# ------------------------------------------------------------------ CRUD


def validate_token(token: str) -> str:
    token = (token or "").strip()
    if not TOKEN_RE.fullmatch(token):
        raise KeyPoolError("That doesn't look like an Apify API token (expected apify_api_ followed by letters/digits)")
    return token


def add_key(db: Session, token: str, label: str | None = None) -> ApifyKey:
    token = validate_token(token)
    h = fingerprint(token)
    if db.scalar(select(ApifyKey.id).where(ApifyKey.token_hash == h)):
        raise KeyPoolError("This key is already in the pool")
    n = (db.scalar(select(func.count(ApifyKey.id))) or 0) + 1
    key = ApifyKey(label=(label or "").strip()[:200] or f"Key {n}", token_enc=encrypt(token), token_hash=h,
                   last4=token[-4:], active=True)
    db.add(key)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise KeyPoolError("This key is already in the pool") from None
    return key


def import_keys(db: Session, filename: str, data: bytes) -> dict:
    """Read every apify_api_… token from an .xlsx or .csv file. A text cell in the same row becomes the label."""
    rows: list[list[str]] = []
    name = filename.lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                rows.append(["" if v is None else str(v).strip() for v in row])
        wb.close()
    elif name.endswith(".csv") or name.endswith(".txt"):
        text = data.decode("utf-8-sig", errors="replace")
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text))]
    else:
        raise KeyPoolError("Upload an .xlsx or .csv file (old .xls files: re-save as .xlsx)")

    added, duplicates, found = 0, 0, 0
    labels: list[str] = []
    for row in rows:
        tokens = [m for cell in row for m in TOKEN_RE.findall(cell)]
        if not tokens:
            continue
        others = [c for c in row if c and not TOKEN_RE.search(c)]
        for i, token in enumerate(tokens):
            found += 1
            label = others[0] if others else None
            if label and len(tokens) > 1:
                label = f"{label} #{i + 1}"
            try:
                k = add_key(db, token, label)
                added += 1
                labels.append(k.label)
            except KeyPoolError:
                duplicates += 1
    if not found:
        raise KeyPoolError("No Apify tokens (apify_api_…) found in the file")
    return {"found": found, "added": added, "duplicates": duplicates, "labels": labels}


def delete_key(db: Session, key_id: int) -> None:
    key = db.get(ApifyKey, key_id)
    if key is None:
        raise KeyPoolError("key not found")
    db.delete(key)  # usage rows keep the label snapshot; FKs are SET NULL
    db.commit()


def update_key(db: Session, key_id: int, label: str | None = None, active: bool | None = None) -> ApifyKey:
    key = db.get(ApifyKey, key_id)
    if key is None:
        raise KeyPoolError("key not found")
    if label is not None and label.strip():
        key.label = label.strip()[:200]
    if active is not None:
        key.active = active
        if active:
            key.exhausted_on = None
            key.last_error = None
    db.commit()
    return key


def log_usage(db: Session, key: ResolvedKey | ApifyKey | None, event: str, run_id: int | None = None,
              detail: str | None = None) -> None:
    key_id = getattr(key, "key_id", None) if isinstance(key, ResolvedKey) else getattr(key, "id", None)
    label = key.label if key is not None else ENV_LABEL
    db.add(ApifyKeyUsage(key_id=key_id, key_label=label, day=today(), at=datetime.now(UTC), event=event,
                         run_id=run_id, detail=redact(detail)[:1000] if detail else None))
    db.commit()


# ------------------------------------------------------------------ daily pick


def _usable(db: Session, d: date):
    return db.scalars(
        select(ApifyKey)
        .where(ApifyKey.active.is_(True), (ApifyKey.exhausted_on.is_(None)) | (ApifyKey.exhausted_on != d))
        .order_by(ApifyKey.last_used_on.is_not(None), ApifyKey.last_used_on, ApifyKey.id)
    ).all()


def _resolved(key: ApifyKey) -> ResolvedKey:
    try:
        return ResolvedKey(key.id, key.label, decrypt(key.token_enc))
    except SecretError as exc:
        raise KeyPoolError(f"{key.label}: {exc}") from None


def has_pool(db: Session) -> bool:
    return bool(db.scalar(select(func.count(ApifyKey.id)).where(ApifyKey.active.is_(True))))


def has_any_token(db: Session) -> bool:
    return has_pool(db) or _env_token() is not None


def current_key(db: Session) -> ResolvedKey | None:
    """Today's key (picking one if none is assigned yet). Falls back to APIFY_TOKEN when the pool is empty."""
    d = today()
    if not has_pool(db):
        tok = _env_token()
        return ResolvedKey(None, ENV_LABEL, tok) if tok else None
    day = db.get(ApifyKeyDay, d)
    if day and day.key_id:
        key = db.get(ApifyKey, day.key_id)
        if key and key.active and key.exhausted_on != d:
            return _resolved(key)
    return _pick(db, d, reason="daily rotation")


def _pick(db: Session, d: date, reason: str) -> ResolvedKey | None:
    candidates = _usable(db, d)
    if not candidates:
        return None
    key = candidates[0]
    key.last_used_on = d
    day = db.get(ApifyKeyDay, d)
    if day is None:
        db.add(ApifyKeyDay(day=d, key_id=key.id, picked_at=datetime.now(UTC)))
    else:
        day.key_id = key.id
        day.picked_at = datetime.now(UTC)
    db.commit()
    log_usage(db, key, "daily_pick", detail=reason)
    return _resolved(key)


def failover(db: Session, key: ResolvedKey, reason: str, run_id: int | None = None) -> ResolvedKey | None:
    """Mark `key` unusable for today and switch today's assignment to the next key."""
    if key.key_id is None:
        return None  # .env token: nothing to switch to
    d = today()
    row = db.get(ApifyKey, key.key_id)
    if row is not None:
        row.exhausted_on = d
        row.last_error = redact(reason)[:500]
        db.commit()
    log_usage(db, key, "failover", run_id=run_id, detail=reason)
    nxt = _pick(db, d, reason=f"failover from {key.label}")
    if nxt is None:
        log.warning("no usable Apify keys left today")
    return nxt


def token_for(db: Session, key_id: int | None) -> ResolvedKey | None:
    """The key that started an existing Apify run (needed to wait for it and read its dataset)."""
    if key_id is None:
        tok = _env_token()
        return ResolvedKey(None, ENV_LABEL, tok) if tok else None
    key = db.get(ApifyKey, key_id)
    return _resolved(key) if key else None


async def key_with_credit(db: Session, needed_usd: float, run_id: int | None = None
                          ) -> tuple[ResolvedKey | None, float | None, str | None]:
    """Today's key, failing over while a key's remaining monthly credit is below `needed_usd`.

    Returns (key, remaining credit or None if unknown, error note)."""
    key = current_key(db)
    tried = 0
    while key is not None and tried < 40:
        tried += 1
        try:
            async with ApifyClient(key.token) as client:
                remaining = summarize_limits(await client.limits())["remaining_usd"]
        except ApifyError as exc:
            if key.key_id is None:
                return key, None, redact(exc)
            key = failover(db, key, f"credit check failed: {exc}", run_id)
            continue
        except Exception as exc:  # noqa: BLE001 - network problems: keep the key, credit unknown
            return key, None, redact(f"{type(exc).__name__}: {exc}")
        if remaining is not None and remaining < needed_usd and key.key_id is not None:
            key = failover(db, key, f"remaining credit ${remaining:.2f} < needed ${needed_usd:.2f}", run_id)
            continue
        return key, remaining, None
    return None, None, "no Apify key with enough credit today" if has_pool(db) else "APIFY_TOKEN not configured"


def summary(db: Session) -> dict:
    d = today()
    keys = db.scalars(select(ApifyKey).order_by(ApifyKey.id)).all()
    uses = dict(db.execute(
        select(ApifyKeyUsage.key_id, func.count()).where(ApifyKeyUsage.event == "run").group_by(ApifyKeyUsage.key_id)
    ).all())
    day = db.get(ApifyKeyDay, d)
    return {
        "today": d,
        "timezone": get_settings().key_rotation_tz,
        "today_key_id": day.key_id if day else None,
        "env_fallback": _env_token() is not None,
        "keys": [key_out(k, uses.get(k.id, 0), day.key_id if day else None) for k in keys],
        "usable_today": len(_usable(db, d)),
    }


def usage(db: Session, limit: int = 200) -> list[dict]:
    rows = db.scalars(select(ApifyKeyUsage).order_by(ApifyKeyUsage.id.desc()).limit(limit)).all()
    return [{"id": r.id, "key_id": r.key_id, "key_label": r.key_label, "day": r.day, "at": r.at, "event": r.event,
             "run_id": r.run_id, "detail": r.detail} for r in rows]
